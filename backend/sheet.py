"""Download the organizer's Google Sheet and expose it as a plain cell grid."""
import io
import urllib.request
from datetime import datetime
from typing import Any, Dict

import openpyxl
from openpyxl.utils import get_column_letter

FORM_SHEET_NAME = "Form Responses 1"

Grid = Dict[int, Dict[str, Any]]  # row number -> {column letter -> value}


def export_url(spreadsheet_id: str) -> str:
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=xlsx"


def download_xlsx(spreadsheet_id: str, timeout: float = 20) -> bytes:
    req = urllib.request.Request(export_url(spreadsheet_id), headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        content = resp.read()
    if not content.startswith(b"PK"):
        raise RuntimeError("Google Sheets did not return an xlsx file (is the sheet still public?)")
    return content


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def parse_form_grid(content: bytes) -> Grid:
    """Return the non-empty cells of the form responses sheet (computed values, not formulas)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb[FORM_SHEET_NAME] if FORM_SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    grid: Grid = {}
    for row in ws.iter_rows():
        cells = {}
        for cell in row:
            value = _clean(cell.value)
            if value is not None:
                cells[get_column_letter(cell.column)] = value
        if cells:
            grid[row[0].row] = cells
    return grid


def as_datetime(value: Any):
    return value if isinstance(value, datetime) else None
